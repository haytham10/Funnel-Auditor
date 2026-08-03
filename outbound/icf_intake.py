"""An ICF Credentialed Coach Finder export -> Leads AND a sourced ICP prefill.

`intake` maps a CSV, and the last three lists were CSVs. This one is an xlsx,
and converting it to a CSV first would silently destroy the column that matters
most: **the load-bearing values are not in the cells.**

`ICF profile` reads the literal string `View profile` in 311 of 311 rows. The
GUID URL that actually reaches the coach's listing lives only in the cell's
hyperlink target. `Email` and `Website` carry hyperlinks too. A converter that
reads cell values produces 311 identical useless strings in a column that looks
fully populated, which is the kind of failure `crm-rows` was written after —
twenty rows went in with no First Name on any of them and the check that passed
them looked at four fields and reported 20/20. So `cell_url` exists, and it
reads the target.

## The second output, and why it is a hint file rather than a verdict

An ICF listing is a form the coach filled in themselves. `Client type`,
`Coaching themes` and `Credential` are the lead's own words about what they sell
and to whom — which is exactly what `docs/spec/02-icp.md` requires for
`sells_to`, "collected from their own words, never inferred". Three ICP fields
arrive answered, for free, on a list where nothing else would answer them
without a fetch.

**But nothing here settles a floor.** The prefill is a separate file, every
value carries `icf_directory` as its source, and no `qualify.Qualification` is
constructed. A directory is stale by construction: a coach who left the UAE two
years ago still reads `Dubai`, and a prefill treated as a fetched fact would
pass the UAE floor confidently and wrongly, which is worse than reading
`unclear`. The activity floor in particular is untouched — it settles on dated
observations, and this file has none.

## Three mappings that are decisions, not transcription

**`Personal and Organizational` means nothing, and it is 79 of the 92 populated
rows.** `qualify.classify_sells_to` already rules that a source claiming both
corporate and individual "says both, so it says nothing" and falls through
rather than letting half a contradiction win. The ICF checkbox is that
contradiction in one cell. Mapping it to `corporates` would put a corporate
identity line in front of 79 people on the strength of a checkbox; mapping it to
`individuals` does the mirror image. It maps to `""`, which draws a generic
line — weaker than an exact match and much stronger than a wrong one.

**`Rate (listed)` is not `top_program_price_aed`.** It is a USD *hourly band*
(`$500-999 per hour`). The ICP field is an AED *program* price. Writing one into
the other is a relabelled number, which is the thing `main.py lint` fails closed
on and the reason `copy/results.csv` exists. It lands in `hourly_rate_usd_band`
under its own name.

**The ICF profile URL is not their site.** It is a directory listing on
`coachingfederation.org`, and putting it in `site_url` would point tier-0 fetch
at ICF's own page for 311 leads and hand every research worker the same
boilerplate. It stays in the prefill, where it is what it is: a re-fetchable
citation for the fields above.

`solo` is not inferred. Nothing on the sheet says it.

openpyxl is imported inside `load_sheet` rather than at module top, so a fresh
clone that has not run `pip install -r requirements.txt` gets an exit-2 message
naming the package instead of an ImportError traceback at CLI parse time.
"""

from __future__ import annotations

import re

from outbound import qualify
from outbound.normalize import Lead, map_row

# The sheet the coaches are on. The workbook also carries README and Summary,
# which are prose and live formulas respectively.
SHEET = "Coaches"

SOURCE = "icf_directory"

# What every prefilled value carries, so a later reader can tell a directory
# self-report from something a worker fetched.
PROVENANCE = "icf_directory"

# ICF's `Client type` checkbox -> `sells_to`. The empty string is a real answer
# here and it is the most common one; see the module docstring.
SELLS_TO = {
    "personal only": "individuals",
    "organizational only": "corporates",
    "personal and organizational": "",
}

# The columns whose values are hyperlinks rather than text.
LINKED_COLUMNS = ("Email", "Website", "ICF profile")

# Phone statuses that say something about the row's freshness. `foreign` is a
# coach listed in the UAE whose number is not — worth a note, never a drop,
# since a UAE-based coach may keep a home-country mobile.
_PHONE_NOTE = {
    "foreign": "ICF phone has a non-UAE country code",
    "invalid": "ICF phone is a placeholder, treat as no phone",
    "review": "ICF phone could not be normalised safely",
}


class ICFIntakeError(RuntimeError):
    """The workbook, the sheet, or openpyxl itself could not be read. Always
    exit 2 — a list nobody could open is not a list of zero coaches."""


def cell_url(cell) -> str:
    """The hyperlink target, falling back to the cell text when the text is
    itself a URL.

    This function is the reason this module exists. `View profile` in 311 of
    311 rows is a column that looks populated and carries nothing, and the only
    place the GUID lives is `cell.hyperlink.target`.

    A `mailto:` target is unwrapped, because the caller wants an address.
    """
    target = ""
    link = getattr(cell, "hyperlink", None)
    if link is not None:
        target = (getattr(link, "target", "") or "").strip()
    if not target:
        text = str(cell.value or "").strip()
        target = text if re.match(r"^(https?://|mailto:|www\.)", text, re.I) else ""
    if target.lower().startswith("mailto:"):
        target = target[7:].split("?", 1)[0]
    return target.strip()


def load_sheet(path: str, *, sheet: str = SHEET) -> list[dict]:
    """One dict per data row, with the three hyperlink targets alongside.

    Every cell arrives as a string. `Profile completeness %` is a float in the
    file and a string everywhere downstream, and a header the alias table has
    never seen must survive to be *reported* as unmapped rather than crashing a
    `.strip()`.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment, not logic
        raise ICFIntakeError(
            f"openpyxl is not installed ({exc}) — it is in requirements.txt, "
            f"and it is needed because this source's real values are cell "
            f"hyperlinks rather than cell text") from exc

    try:
        book = openpyxl.load_workbook(path)
    except Exception as exc:
        raise ICFIntakeError(f"could not open {path} ({type(exc).__name__}: {exc})") from exc

    if sheet not in book.sheetnames:
        raise ICFIntakeError(
            f"{path} has no sheet named {sheet!r} — it has "
            f"{', '.join(book.sheetnames)}")

    grid = book[sheet]
    rows = grid.iter_rows()
    try:
        header_cells = next(rows)
    except StopIteration:
        raise ICFIntakeError(f"{path}:{sheet} is empty") from None

    headers = [str(c.value or "").strip() for c in header_cells]
    linked = {h: i for i, h in enumerate(headers) if h in LINKED_COLUMNS}

    out = []
    for cells in rows:
        if not any(c.value not in (None, "") for c in cells):
            continue
        record = {h: ("" if c.value is None else str(c.value).strip())
                  for h, c in zip(headers, cells) if h}
        for header, index in linked.items():
            if index < len(cells):
                record[f"_{header.lower().replace(' ', '_')}_url"] = cell_url(cells[index])
        out.append(record)
    return out


def to_lead(record: dict, *, source: str = SOURCE) -> Lead:
    """One ICF row -> a Lead, built through `normalize.map_row`.

    Routed through `map_row` rather than constructing a `Lead` directly so the
    name split, the slug, `classify_site` on the website and the free
    handle-ownership note are one implementation and not two — the same reason
    `ig_intake.to_lead` does it.

    The address and the website come from the hyperlink targets when they exist:
    a `Website` cell reads `http://www.example.com` and its target is the
    normalised `http://www.example.com/`, and where the two disagree the target
    is the one a browser would actually follow.
    """
    email = record.get("_email_url") or record.get("Email", "")
    website = record.get("_website_url") or record.get("Website", "")

    # The credential and the themes together are the closest thing this source
    # has to a positioning line, and `headline` is what `classify_coach_type`
    # and the SERP query both read.
    headline = " ".join(p for p in (record.get("Credential", ""),
                                    record.get("Coaching themes", "")) if p)

    lead = map_row({
        "name": record.get("Name", ""),
        "email": email,
        "website": website,
        "city": record.get("City", ""),
        "phone": record.get("Phone (E.164)", ""),
        "headline": headline,
    }, source=source)

    lead.country = "United Arab Emirates"

    note = _PHONE_NOTE.get((record.get("Phone status") or "").strip().lower())
    if note:
        lead.notes.append(note)
    return lead


def prefill(record: dict) -> dict:
    """The ICF columns as ICP fields, every value carrying where it came from.

    `coach_type` goes through `qualify.classify_coach_type` rather than a second
    keyword table — the patterns there are tested and already encode that
    LinkedIn wins a disagreement, which matters because `research` will later
    read a real LinkedIn headline for these same leads and must be able to
    overrule a directory checkbox.

    `sells_to` is mapped from the enum directly rather than run through
    `classify_sells_to`, because `Client type` is a controlled vocabulary and
    not prose. The RULE is that function's — a source saying both says nothing —
    and `SELLS_TO` is that rule written for this source's three values.
    """
    themes = record.get("Coaching themes", "")
    coach_type, _ = qualify.classify_coach_type(site_text=themes)

    client_type = (record.get("Client type") or "").strip().lower()
    sells_to = SELLS_TO.get(client_type, "")

    out = {
        "name": record.get("Name", ""),
        "credential": record.get("Credential", ""),
        "city": record.get("City", ""),
        "emirate": record.get("Emirate", ""),
        "coach_type": coach_type,
        "coach_type_source": f"{PROVENANCE} (coaching themes)" if coach_type else "",
        "sells_to": sells_to,
        "sells_to_source": _sells_to_source(client_type, sells_to),
        # NOT top_program_price_aed. A USD hourly band is not an AED program
        # price, and relabelling one as the other is what `lint` fails on.
        "hourly_rate_usd_band": record.get("Rate (listed)", ""),
        "fee_range_usd": record.get("Fee range", ""),
        "icf_profile_url": record.get("_icf_profile_url", ""),
        "icf_key": record.get("ICF key", ""),
        "profile_completeness": record.get("Profile completeness %", ""),
        "languages": record.get("Languages", ""),
        "phone": record.get("Phone (E.164)", ""),
        "phone_status": record.get("Phone status", ""),
        "coaching_themes": themes,
        "client_type": record.get("Client type", ""),
        "source": PROVENANCE,
    }
    return out


def _sells_to_source(client_type: str, sells_to: str) -> str:
    """Why `sells_to` says what it says — including when it says nothing.

    An empty field with no reason attached is indistinguishable from a field
    nobody looked at, and 219 of 311 rows are empty because the coach left the
    box blank while 79 are empty because the coach ticked both.
    """
    if not client_type:
        return f"{PROVENANCE} (client type blank on the listing)"
    if not sells_to:
        return (f"{PROVENANCE} (client type {client_type!r} — says both, "
                f"so it says nothing)")
    return f"{PROVENANCE} (client type {client_type!r})"


def ingest(path: str, *, sheet: str = SHEET,
           source: str = SOURCE) -> tuple[list[Lead], dict]:
    """(leads, prefill keyed by `fetch.lead_key`).

    Keyed on `fetch.lead_key` because that is the one definition of "which lead
    is this" — the ledger, `resolve`, `observe` and `select` all join on it, and
    a second spelling would join to nothing.
    """
    from outbound.fetch import lead_key

    records = load_sheet(path, sheet=sheet)
    leads, prefills = [], {}
    for record in records:
        lead = to_lead(record, source=source)
        leads.append(lead)
        prefills[lead_key(lead)] = prefill(record)
    return leads, prefills


def unmapped_columns(path: str, *, sheet: str = SHEET) -> list[str]:
    """Headers the alias table does not map, reported rather than dropped.

    `intake` reports these for a CSV and the reason is the same here: a column
    nobody mapped is either a field this machine should learn or a field it has
    decided to ignore, and silence cannot tell the two apart.
    """
    from outbound.normalize import COLUMN_ALIASES, _canon_header

    records = load_sheet(path, sheet=sheet)
    if not records:
        return []
    headers = [h for h in records[0] if not h.startswith("_")]
    return [h for h in headers if _canon_header(h) not in COLUMN_ALIASES]


# The columns enrichment adds, in order, after the source sheet's own 31.
# `blank means not established` is the contract, and the two note columns are
# what make a blank readable: which KIND of blank it is.
ENRICHED_COLUMNS = (
    "LinkedIn URL", "LinkedIn source", "LinkedIn evidence",
    "Instagram URL", "Instagram source",
    "Website (found)", "Website source",
    "Channel verdict", "Channel note",
    "Email status", "Email result", "Email free-mail", "Email note",
    "Coach type (ICF)", "Sells to (ICF)", "Hourly rate band (ICF)",
    "ICF profile URL",
    "Dedupe", "Enriched at", "Batch",
)


def enriched_row(lead: Lead, *, prefill: dict, channels: dict, verify: dict,
                 contacted: bool, batch: str, at: str) -> dict:
    """One lead's enrichment columns.

    **A blank is never bare.** `Channel verdict` and `Channel note` say which of
    the four blanks this is — nothing found, found and rejected as somebody
    else's, two people of that name and nothing separating them, or not searched
    at all. An operator deciding whether to spend thirty seconds looking by hand
    needs that distinction, and it is the whole reason the verdicts are not
    collapsed into a boolean.
    """
    found = (channels or {}).get("accepted") or {}
    verdict = (channels or {}).get("verdict", "")

    # Provenance is derived from the lead in hand rather than read off the
    # channel file. The file's own `from_row` is written by a command that once
    # dropped the key entirely, and the result was 84 rows whose "found website"
    # was a copy of the website the sheet already listed — a column that looked
    # like 117 discoveries and contained 33. The lead is the authority on what
    # the row already carried, and it is right here.
    row_values = {"linkedin_url": lead.linkedin_url,
                  "instagram_url": lead.instagram_url,
                  "site_url": lead.site_url}

    def source(field: str) -> str:
        value = found.get(field)
        if not value:
            return ""
        return "ICF row" if value == row_values.get(field) else "search"

    def discovered(field: str) -> str:
        """Only what the row did not already have."""
        value = found.get(field, "")
        return "" if value == row_values.get(field) else value

    evidence = ""
    for cand in (channels or {}).get("candidates") or []:
        if cand.get("platform") == "linkedin" and cand.get("url") == found.get("linkedin_url"):
            evidence = f"{cand.get('tier', '')}: {cand.get('why', '')}"
            break

    domain = lead.email.split("@")[-1].lower() if lead.email else ""
    free = domain in _FREE_MAIL

    return {
        "LinkedIn URL": found.get("linkedin_url", ""),
        "LinkedIn source": source("linkedin_url"),
        "LinkedIn evidence": evidence,
        "Instagram URL": found.get("instagram_url", ""),
        "Instagram source": source("instagram_url"),
        # Never overwrites the sheet's own Website column — a found site is a
        # different claim from a listed one and lives in its own column, and a
        # copy of the listed one is not a discovery at all.
        "Website (found)": discovered("site_url"),
        "Website source": "search" if discovered("site_url") else "",
        "Channel verdict": verdict or "NOT SEARCHED",
        "Channel note": (channels or {}).get("reason", "") or
        "excluded before any paid call — see Dedupe",
        "Email status": (verify or {}).get("status", ""),
        "Email result": (verify or {}).get("result", ""),
        "Email free-mail": "yes" if free else "no",
        "Email note": _email_note(verify or {}, free),
        "Coach type (ICF)": (prefill or {}).get("coach_type", ""),
        "Sells to (ICF)": (prefill or {}).get("sells_to", ""),
        "Hourly rate band (ICF)": (prefill or {}).get("hourly_rate_usd_band", ""),
        "ICF profile URL": (prefill or {}).get("icf_profile_url", ""),
        "Dedupe": "ALREADY CONTACTED — do not send" if contacted else "clear",
        "Enriched at": at,
        "Batch": batch,
    }


_FREE_MAIL = {"gmail.com", "hotmail.com", "yahoo.com", "live.com", "outlook.com",
              "icloud.com", "me.com", "aol.com", "gmx.net", "hotmail.co.uk",
              "yahoo.co.uk", "googlemail.com", "yahoo.fr", "msn.com",
              "hotmail.fr", "mail.ru", "protonmail.com", "yandex.com"}


def _email_note(verify: dict, free: bool) -> str:
    """What the verdict does and does not prove.

    A `catch_all` on a branded domain is the common case here and it is not a
    weak PASS — the server accepts every address, so nothing was learned about
    this mailbox. And a clean `valid` on a directory address proves the mailbox
    exists, never that the coach still reads it.
    """
    result = verify.get("result", "")
    if not result:
        return "not verified"
    if result == "catch_all":
        return ("the domain accepts every address, so this mailbox was neither "
                "confirmed nor ruled out")
    if result == "invalid":
        return "the mailbox does not exist — a stale directory row"
    if result == "unknown":
        return "the verifier could not settle it"
    if result == "valid" and free:
        return ("mailbox exists. A free-mail address on a directory listing "
                "proves deliverability, never that it is still read")
    if result == "valid":
        return "mailbox exists on their own domain"
    return result


def write_enriched(src_path: str, rows: list[dict], out_path: str, *,
                   sheet: str = SHEET) -> str:
    """The source workbook with the enrichment columns appended.

    The original 31 columns are left exactly as they arrived, in order,
    hyperlinks and all — this is the file Haytham already knows how to read, and
    a rewrite that "tidied" it would make every row need re-checking. `rows` is
    keyed by the row's ICF key, which is the only identifier the sheet carries
    that survives a name normalisation.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment, not logic
        raise ICFIntakeError(f"openpyxl is not installed ({exc})") from exc

    try:
        book = openpyxl.load_workbook(src_path)
    except Exception as exc:
        raise ICFIntakeError(f"could not open {src_path} "
                             f"({type(exc).__name__}: {exc})") from exc
    if sheet not in book.sheetnames:
        raise ICFIntakeError(f"{src_path} has no sheet named {sheet!r}")

    grid = book[sheet]
    headers = [str(c.value or "").strip() for c in next(grid.iter_rows())]
    try:
        key_column = headers.index("ICF key") + 1
    except ValueError:
        raise ICFIntakeError(f"{src_path}:{sheet} has no 'ICF key' column, so "
                             f"enrichment cannot be joined back to it") from None

    start = len(headers) + 1
    for offset, column in enumerate(ENRICHED_COLUMNS):
        grid.cell(1, start + offset, column)

    written = 0
    for index in range(2, grid.max_row + 1):
        key = str(grid.cell(index, key_column).value or "").strip()
        row = rows.get(key) if isinstance(rows, dict) else None
        if not row:
            continue
        written += 1
        for offset, column in enumerate(ENRICHED_COLUMNS):
            grid.cell(index, start + offset, row.get(column, ""))

    if not written:
        raise ICFIntakeError("no row joined on 'ICF key' — the enrichment and "
                             "the workbook do not describe the same list")
    book.save(out_path)
    return out_path


def report(leads: list[Lead], prefills: dict) -> str:
    """The quotable profile, in `intake`'s voice.

    The free-mail share is here because it is the single most consequential
    fact about this list after the address count: `email_enrich` refuses free
    provider domains outright, so the no-address fallback has no path at all on
    those leads, and a SERP-found address on one can never be corroborated by
    `on_lead_domain`.
    """
    from audit.email_enrich import FREE_PROVIDERS
    from audit.urls import registrable_domain

    total = len(leads)
    with_email = sum(1 for l in leads if l.email)
    with_site = sum(1 for l in leads if l.site_url)
    free_mail = sum(1 for l in leads
                    if l.email and registrable_domain("http://" + l.email.split("@")[-1])
                    in FREE_PROVIDERS)
    no_target = sum(1 for l in leads if not l.has_research_target())
    nameless = [l for l in leads if not l.name]

    filled = {}
    for field_name in ("coach_type", "sells_to", "icf_profile_url", "credential"):
        filled[field_name] = sum(1 for p in prefills.values() if p.get(field_name))

    lines = [
        f"ICF INTAKE: {total} lead(s) from {SOURCE}",
        f"  {with_email}/{total} with an email, {free_mail} of them free-mail "
        f"— `email-enrich` has no path on those, it refuses free providers",
        f"  {with_site}/{total} with a live site; {no_target} have no research "
        f"target at all beyond their name",
        f"  prefill: coach_type {filled['coach_type']}/{total}, "
        f"sells_to {filled['sells_to']}/{total}, "
        f"credential {filled['credential']}/{total}, "
        f"ICF profile URL {filled['icf_profile_url']}/{total}",
        "  The prefill is the coach's own words on a directory form. It is a "
        "hint file — nothing in it settles a floor, and the activity floor is "
        "untouched.",
    ]
    for lead in nameless:
        lines.append(f"  NO NAME  {lead.email or lead.site_url or '(empty row)'} "
                     f"— `dedupe` cannot check a nameless row against the wall")
    return "\n".join(lines)
